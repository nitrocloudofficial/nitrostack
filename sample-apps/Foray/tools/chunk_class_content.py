"""
FORAY — turn the uploaded class content (sem V.zip) into MaterialChunks.

    python tools/chunk_class_content.py

The textbook is already chunked by tools/chunk_material.py. This adds the
REST of the class's prescribed material from the uploaded semester archive:

    notes/sem V/Computer Networks/PPTS/*.pptx     lecture slide decks
    notes/sem V/Computer Networks/PPTS/*.pdf      problem sets
    notes/sem V/Computer Networks/Wireshark_*.pdf lab handouts

Same rules as the textbook chunker:

  * tagging is KEYWORD-BASED AND DETERMINISTIC — the same hand-authored
    CONCEPT_KEYWORDS table, no embeddings (§9)
  * untagged chunks are dropped, not guessed at
  * `page` is what a student would turn to: the SLIDE NUMBER for decks,
    the printed page for PDFs

Ids continue the MAT_ namespace from 5001 so the two corpora never collide.
Output: data/materials/class_chunks.json — loaded by api/db.load_static
alongside chunks.json.
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import fitz

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from tools.chunk_material import CONCEPT_KEYWORDS, _matches, clean  # noqa: E402

CONTENT = ROOT / "notes" / "sem V" / "Computer Networks"
OUT = ROOT / "data" / "materials" / "class_chunks.json"

# Slides are terse; a 120-char slide can still carry one precise idea.
MIN_SLIDE_CHARS = 120
MIN_PDF_CHARS = 240
MAX_CHARS = 1400

DECKS = [
    ("PPTS/Transport Layer.pptx", "Transport Layer slides (CS3591)"),
    ("PPTS/Application Layer.pptx", "Application Layer slides (CS3591)"),
    ("PPTS/Intro.pptx", "Introduction slides (CS3591)"),
]
PDFS = [
    ("PPTS/Chapter3-Problems.pdf", "Ch.3 problem set (CS3591)"),
    ("PPTS/Chapter2-Problems.pdf", "Ch.2 problem set (CS3591)"),
    ("Wireshark_Intro_v8.0.pdf", "Wireshark intro lab (CS3591)"),
    ("Wireshark_HTTP_v8.0.pdf", "Wireshark HTTP lab (CS3591)"),
]


def tag(text: str) -> list[str]:
    low = text.lower()
    return [c for c, kws in CONCEPT_KEYWORDS.items()
            if any(_matches(k, low) for k in kws)]


def slide_text(slide) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs).strip()
                if line:
                    parts.append(line)
    return clean("\n".join(parts))


def chunk_deck(path: Path, source: str, start_id: int) -> tuple[list[dict], int]:
    from pptx import Presentation

    chunks: list[dict] = []
    n = start_id
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, start=1):
        body = slide_text(slide)
        if len(body) < MIN_SLIDE_CHARS:
            continue
        cids = tag(body)
        if not cids:
            continue
        n += 1
        chunks.append({
            "id": f"MAT_{n:04d}",
            "source": source,
            "page": i,                      # the SLIDE number
            "text": body[:MAX_CHARS],
            "concept_ids": cids,
        })
    return chunks, n


def chunk_pdf(path: Path, source: str, start_id: int) -> tuple[list[dict], int]:
    chunks: list[dict] = []
    n = start_id
    doc = fitz.open(path)
    for i in range(doc.page_count):
        body = clean(doc[i].get_text())
        if len(body) < MIN_PDF_CHARS:
            continue
        cids = tag(body)
        if not cids:
            continue
        n += 1
        chunks.append({
            "id": f"MAT_{n:04d}",
            "source": source,
            "page": i + 1,                  # standalone doc: printed == index+1
            "text": body[:MAX_CHARS],
            "concept_ids": cids,
        })
    return chunks, n


def main() -> None:
    if not CONTENT.exists():
        raise SystemExit(f"class content not found: {CONTENT} — "
                         f"extract sem V.zip under notes/ first")

    all_chunks: list[dict] = []
    n = 5000                                # MAT_5001 onward: never collides
    for rel, source in DECKS:
        p = CONTENT / rel
        if not p.exists():
            print(f"  skip (missing): {rel}")
            continue
        got, n = chunk_deck(p, source, n)
        all_chunks.extend(got)
        print(f"  {source:38} {len(got):4d} chunks")
    for rel, source in PDFS:
        p = CONTENT / rel
        if not p.exists():
            print(f"  skip (missing): {rel}")
            continue
        got, n = chunk_pdf(p, source, n)
        all_chunks.extend(got)
        print(f"  {source:38} {len(got):4d} chunks")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(all_chunks, indent=1, ensure_ascii=False),
                   encoding="utf-8")

    per: dict[str, int] = {}
    for ch in all_chunks:
        for c in ch["concept_ids"]:
            per[c] = per.get(c, 0) + 1
    print(f"\nwrote {len(all_chunks)} chunks -> {OUT.relative_to(ROOT)}")
    print("chunks per concept:")
    for c in sorted(per):
        print(f"  {c}  {per[c]:4d}  " + "#" * min(40, per[c]))


if __name__ == "__main__":
    main()
