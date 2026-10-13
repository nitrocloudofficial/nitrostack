import os
import logging
from typing import Dict, Any, List
from PIL import Image
import PyPDF2
from docx import Document as DocxDocument
from pptx import Presentation

logger = logging.getLogger(__name__)

try:
    import pytesseract
except ImportError:
    pytesseract = None


class FileProcessor:
    """Utility class for extracting text and metadata from various file formats."""

    @staticmethod
    def extract_pdf(file_path: str) -> Dict[str, Any]:
        """Extract text page-by-page from a PDF file using PyPDF2."""
        pages_content: List[Dict[str, Any]] = []
        full_text = []

        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            page_count = len(reader.pages)
            
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                text = text.strip()
                if text:
                    full_text.append(text)
                pages_content.append({"page": idx + 1, "text": text})

        return {
            "text": "\n\n".join(full_text),
            "page_count": page_count,
            "pages": pages_content
        }

    @staticmethod
    def extract_docx(file_path: str) -> Dict[str, Any]:
        """Extract text from Microsoft Word (.docx) documents."""
        doc = DocxDocument(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        # Approximate page count (~500 words per page)
        word_count = len(full_text.split())
        estimated_pages = max(1, (word_count + 499) // 500)

        return {
            "text": full_text,
            "page_count": estimated_pages,
            "pages": [{"page": 1, "text": full_text}]
        }

    @staticmethod
    def extract_pptx(file_path: str) -> Dict[str, Any]:
        """Extract text slide-by-slide from Microsoft PowerPoint (.pptx) presentations."""
        prs = Presentation(file_path)
        slide_count = len(prs.slides)
        pages_content: List[Dict[str, Any]] = []
        full_text = []

        for idx, slide in enumerate(prs.slides):
            slide_text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text_parts.append(shape.text.strip())
            
            combined_slide_text = "\n".join(slide_text_parts)
            if combined_slide_text:
                full_text.append(f"--- Slide {idx + 1} ---\n{combined_slide_text}")
            pages_content.append({"page": idx + 1, "text": combined_slide_text})

        return {
            "text": "\n\n".join(full_text),
            "page_count": slide_count,
            "pages": pages_content
        }

    @staticmethod
    def extract_image(file_path: str) -> Dict[str, Any]:
        """Perform OCR on an image file using Pillow and pytesseract."""
        if pytesseract is None:
            logger.warning("pytesseract is not installed. OCR skipped.")
            return {"text": "", "page_count": 1, "pages": [{"page": 1, "text": ""}]}

        try:
            image = Image.open(file_path)
            extracted_text = pytesseract.image_to_string(image).strip()
            return {
                "text": extracted_text,
                "page_count": 1,
                "pages": [{"page": 1, "text": extracted_text}]
            }
        except Exception as e:
            logger.error(f"Error performing OCR on {file_path}: {e}")
            return {"text": "", "page_count": 1, "pages": [{"page": 1, "text": ""}]}

    @staticmethod
    def extract_plain_text(file_path: str) -> Dict[str, Any]:
        """Extract text from plain text or markdown files."""
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read().strip()

        word_count = len(text.split())
        estimated_pages = max(1, (word_count + 499) // 500)

        return {
            "text": text,
            "page_count": estimated_pages,
            "pages": [{"page": 1, "text": text}]
        }

    @classmethod
    def process_file(cls, file_path: str) -> Dict[str, Any]:
        """
        Process a file based on its extension and extract text content and page count.
        
        Args:
            file_path: Absolute or relative path to the file.
            
        Returns:
            Dict containing 'text', 'page_count', and 'pages'.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".pdf":
            return cls.extract_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls.extract_docx(file_path)
        elif ext in [".pptx", ".ppt"]:
            return cls.extract_pptx(file_path)
        elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"]:
            return cls.extract_image(file_path)
        elif ext in [".txt", ".md", ".csv", ".json", ".log"]:
            return cls.extract_plain_text(file_path)
        else:
            # Fallback to plain text reading
            try:
                return cls.extract_plain_text(file_path)
            except Exception as e:
                logger.error(f"Failed to process unsupported file type {ext}: {e}")
                raise ValueError(f"Unsupported file type: {ext}")
